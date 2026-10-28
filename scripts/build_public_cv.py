"""Build the reviewed tourism CV Word source; render and review before replacing PDF files."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
BASE=Path(__file__).resolve().parents[1]
OUT=BASE/'documents'
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(16);sec.bottom_margin=Mm(16);sec.left_margin=Mm(19);sec.right_margin=Mm(19)
sec.header_distance=Mm(7);sec.footer_distance=Mm(8)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
 s=doc.styles[n];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if s.element.get_or_add_rPr().find(qn('w:rFonts')) is None else None
 s.paragraph_format.widow_control=True
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.07
for n in ['Heading 1','Heading 2']:
 s=doc.styles[n];s.font.size=Pt(11.5 if n=='Heading 1' else 10.5);s.font.bold=True
 s.paragraph_format.space_before=Pt(10 if n=='Heading 1' else 7);s.paragraph_format.space_after=Pt(4);s.paragraph_format.keep_with_next=True
style=doc.styles['Title'];style.font.name='Georgia';style.font.size=Pt(27);style.paragraph_format.space_after=Pt(4)
style=doc.styles['List Bullet'];style.font.size=Pt(10.5);style.paragraph_format.space_after=Pt(3);style.paragraph_format.line_spacing=1.05
style.paragraph_format.left_indent=Mm(3.5);style.paragraph_format.first_line_indent=Mm(-3.5)
# Use explicit real numbering for simple extraction and keyboard/screen-reader navigation.
numroot=doc.part.numbering_part.element
abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'42')
lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
for tag,val in [('start','1'),('numFmt','bullet'),('lvlText','•'),('lvlJc','left')]:
 e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lvl.append(e)
ppr=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'200');ind.set(qn('w:hanging'),'200');ppr.append(ind);lvl.append(ppr);abstract.append(lvl);numroot.append(abstract)
num=OxmlElement('w:num');num.set(qn('w:numId'),'42');aid=OxmlElement('w:abstractNumId');aid.set(qn('w:val'),'42');num.append(aid);numroot.append(num)
def p(text='',style=None,bold=None):
 a=doc.add_paragraph(style=style);r=a.add_run(text);r.bold=bold
 if style=='Title':r.font.name='Georgia'
 return a
def link(par,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'253D48');pr.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def bullet(text):
 a=p(text,'List Bullet');pr=a._p.get_or_add_pPr();np=OxmlElement('w:numPr');il=OxmlElement('w:ilvl');il.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'42');np.append(il);np.append(ni);pr.append(np)
def employer(name,location,function,dates,items):
 p(name+' | '+location,'Heading 2')
 a=p(function+' | '+dates);a.paragraph_format.keep_with_next=True
 for r in a.runs:r.font.size=Pt(9.5);r.font.color.rgb=RGBColor.from_string('444444')
 for text in items:bullet(text)
def heading(text):p(text,'Heading 1')
p('Ahmed Mahmoud','Title')
p('Travel and Tourism Operations',bold=True)
a=p('Destination services · Tourism product support · MICE and event coordination');a.paragraph_format.space_after=Pt(6)
a=p('Den Haag, Netherlands · +31 6 2889 5967');a.paragraph_format.space_after=Pt(2)
a=p();link(a,'ahmed.mahmoud.nl.work@gmail.com','mailto:ahmed.mahmoud.nl.work@gmail.com');a.add_run(' · ');link(a,'ahmedqualityops.com','https://ahmedqualityops.com/');a.paragraph_format.space_after=Pt(2)
a=p();link(a,'linkedin.com/in/ahmed-mahmoud-tourism','https://www.linkedin.com/in/ahmed-mahmoud-tourism/');a.paragraph_format.space_after=Pt(2)
a=p('Authorised to work in the Netherlands without employer sponsorship.');a.paragraph_format.space_after=Pt(8)
heading('Professional profile')
p('Travel and tourism professional with experience across Egypt and the UAE in destination services, B2B reservations, travel arrangements and cross-functional operations support. Coordinated guests, suppliers, transport and operational handovers. Based in Den Haag and seeking a role in tourism operations, tourism product support or MICE and event coordination.')
heading('Relevant capabilities')
p('Guest movements and transfers · Supplier coordination · Itineraries and quotations · B2B reservations · MICE and group logistics · Service recovery · Procedures and handovers · Process improvement')
heading('Professional experience')
employer('Masarra DMC / Miral Experiences','Abu Dhabi, UAE','Operations Department and cross-functional support','Dec 2024 - Oct 2026',[
'Supported guest communication, transport, ticketing, suppliers and service recovery as part of the operations team.',
'Contributed to MICE, familiarisation trips and industry events through airport and hotel representation, guest movements and stakeholder updates.',
'Contributed to airport-representative procedures, colleague coaching, Arabic CMS localisation and a Travel Desk improvement project.'
])
employer('Oceanair Travels','Dubai, UAE','Destination operations and guest coordination','Sep 2022 - Nov 2024',[
'Coordinated destination programmes, hotel pickups, transport, tickets, attraction timing and supplier handovers.',
'Handled guest communication and service recovery during delays and programme changes.'
])
employer('WebBeds Middle East and Africa','Cairo, Egypt','B2B reservations and partner support','Jun 2022 - Sep 2022',[
'Handled booking cases using Juniper and Freshdesk, including amendments, discrepancies, refunds and relocations.',
'Supported French-speaking hotel markets in Tunisia, Morocco and Algeria, working with partners and internal teams.'
])
employer('Fun and Sun / TUI CIS','Sharm El Sheikh, Egypt','Excursions, airport transfers and guest services','Jul 2021 - Apr 2022',[
'Coordinated excursions, arrivals, departures and transfers; supported guest issues and programme changes, including CIS and Russian-speaking markets.'
])
doc.add_page_break()
heading('Earlier professional experience')
employer('Alnatalie Red Sea Tours','Cairo, Egypt','Transport and fleet coordination','Jan 2020 - Jun 2021',[
'Coordinated vehicle and driver allocation, GPS follow-up, maintenance scheduling and operational readiness.',
'Supported transport for the 2021 IHF Men\'s World Championship, including airport, hotel and competition-venue movements.'
])
employer('Elmadina Tours','Cairo, Egypt','Travel consulting and package arrangements','Jan 2013 - Jan 2020',[
'Designed, quoted and sold leisure and honeymoon packages; coordinated itineraries, accommodation, suppliers and customer follow-up.',
'Coordinated Hajj and Umrah arrangements, flights, visas and accommodation using SAFA and Amadeus workflows.'
])
heading('Selected projects')
a=p();a.add_run('Travel Desk improvement').bold=True;a.add_run(' | Workplace project')
p('Applied DMAIC to agent allocation, training, multilingual materials and control planning. The financial opportunity was a conditional projection, not realised revenue.')
a=p();a.add_run('Amsterdam tourism product discovery').bold=True;a.add_run(' | Portfolio research, pre-pilot')
p('Developed research and planning materials covering customer evidence, product requirements, supplier feasibility, unit economics and pilot controls. No commercial departures or measured post-pilot outcomes are claimed.')
a=p();a.add_run('InfraQuote and InfraDispatch').bold=True;a.add_run(' | Browser demonstrations')
p('Built portfolio demonstrations of quotation review, movement planning and operational handovers using sample or operator-entered data. These are prototypes, not production client deployments.')
heading('Selected qualifications')
for text in [
'IATA Travel Operations diploma | Mar 2026',
'IATA MICE and Sales diploma | Jul 2025',
'IATA Tour Production diploma | Sep 2024',
'IATA Travel and Tourism Consultant diploma | May 2021',
'Certified Associate in Project Management (CAPM), PMI | Apr 2024 - Apr 2027',
'Lean Six Sigma Black Belt training certificate, Anexas Europe / Miral LLC | Jul 2025',
'ITIL V4 Foundation | Recorded certificate dates Aug 2024 - Aug 2027']:
 a=p(text);a.paragraph_format.space_after=Pt(2);a.paragraph_format.line_spacing=1
heading('Education')
p('Licence / Bachelor of Law | Faculty of Law, Cairo University | 2006 - 2010')
heading('Languages')
p('Arabic: native · English: advanced professional · French: B1+ (self-assessed), DELF B1 · Russian: A2 · Dutch: A1, actively studying')
heading('Systems')
p('Juniper · Freshdesk · Amadeus · SAFA · BOS · Otolo Homes · Strapi CMS · Excel and Microsoft Office · Google Workspace')
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('Ahmed Mahmoud  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.core_properties.author='Ahmed Mahmoud';doc.core_properties.title='Ahmed Mahmoud Travel and Tourism Operations CV';doc.core_properties.subject='Tourism operations, product support and MICE coordination';doc.core_properties.comments=''
# English document language.
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-GB');normal.element.get_or_add_rPr().append(lang)
# Strip inherited title rules and theme font overrides.
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for fonts in st.element.iter(qn('w:rFonts')):
  for key in list(fonts.attrib):
   if 'Theme' in key or key.endswith('cstheme'):del fonts.attrib[key]
doc.save(OUT/'Ahmed-Mahmoud-CV.docx')
print(OUT/'Ahmed-Mahmoud-CV.docx')
